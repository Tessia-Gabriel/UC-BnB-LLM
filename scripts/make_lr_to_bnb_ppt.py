#!/usr/bin/env python3
"""Generate the 7-slide LR-to-BnB project presentation."""

from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables/UC_LR_to_BnB_LLM4Branch_Progress_20260921.pptx"

W, H = 13.333, 7.5
FONT = "Noto Sans CJK SC"
BG = "F7F9FC"
WHITE = "FFFFFF"
NAVY = "10223E"
BLUE = "3568D4"
TEAL = "0F9D99"
ORANGE = "E9A23B"
RED = "D75B5B"
TEXT = "1F2D3D"
MUTED = "667085"
LINE = "DCE3ED"
PALE_BLUE = "EAF0FC"
PALE_TEAL = "E6F5F3"
PALE_ORANGE = "FCF2DF"
PALE_RED = "FBEAEA"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def add_text(slide, text, x, y, w, h, size=18, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.04,
             font=FONT, line_spacing=1.05):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def rect(slide, x, y, w, h, fill=WHITE, line=LINE, radius=True, width=1):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line); shape.line.width = Pt(width)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def line(slide, x1, y1, x2, y2, color=LINE, width=1.5, dash=None):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = rgb(color); shape.line.width = Pt(width)
    if dash is not None:
        shape.line.dash_style = dash
    return shape


def circle_label(slide, text, x, y, diameter=0.42, fill=BLUE, size=14):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(diameter), Inches(diameter))
    c.fill.solid(); c.fill.fore_color.rgb = rgb(fill); c.line.fill.background()
    add_text(slide, text, x, y - 0.005, diameter, diameter, size=size, color=WHITE,
             bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)
    return c


def title(slide, kicker, heading, page):
    add_text(slide, kicker.upper(), 0.72, 0.31, 5.6, 0.28, 10, TEAL, True)
    add_text(slide, heading, 0.72, 0.62, 11.9, 0.55, 27, NAVY, True)
    line(slide, 0.72, 1.25, 12.62, 1.25, LINE, 1)
    add_text(slide, f"{page:02d}", 12.35, 0.34, 0.3, 0.22, 9, MUTED,
             align=PP_ALIGN.RIGHT)


def footer(slide, source):
    line(slide, 0.72, 7.08, 12.62, 7.08, LINE, 0.8)
    add_text(slide, source, 0.72, 7.13, 11.7, 0.2, 8, MUTED)


def bullet(slide, text, x, y, w, color=TEXT, size=16, dot=TEAL, bold_prefix=None):
    circle_label(slide, "", x, y + 0.09, 0.12, dot, 1)
    if bold_prefix and text.startswith(bold_prefix):
        box = slide.shapes.add_textbox(Inches(x + 0.23), Inches(y), Inches(w - 0.23), Inches(0.55))
        tf = box.text_frame; tf.clear(); tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        r = p.add_run(); r.text = bold_prefix; r.font.name = FONT; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = rgb(color)
        r = p.add_run(); r.text = text[len(bold_prefix):]; r.font.name = FONT; r.font.size = Pt(size); r.font.color.rgb = rgb(color)
        return box
    return add_text(slide, text, x + 0.23, y, w - 0.23, 0.55, size, color)


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    # 1 — title
    s = prs.slides.add_slide(blank)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(NAVY)
    s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(H)).fill.solid()
    s.shapes[-1].fill.fore_color.rgb = rgb(TEAL); s.shapes[-1].line.fill.background()
    add_text(s, "UNIT COMMITMENT · RESEARCH UPDATE", 0.86, 0.66, 7.3, 0.28, 11, "77D5D0", True)
    add_text(s, "从拉格朗日松弛到\n学习增强的分支定界", 0.86, 1.22, 8.3, 1.65, 31, WHITE, True, line_spacing=0.96)
    add_text(s, "路线转向的原因 · LLM4Branch 研究进展 · 当前复现证据", 0.89, 3.06, 8.7, 0.42, 16, "C8D3E3")
    rect(s, 9.54, 1.18, 2.78, 4.55, fill="162F53", line="294667")
    circle_label(s, "LR", 10.42, 1.75, 1.02, ORANGE, 20)
    line(s, 10.93, 2.82, 10.93, 3.52, "7590B2", 2)
    tri = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(10.68), Inches(3.27), Inches(0.5), Inches(0.55))
    tri.fill.solid(); tri.fill.fore_color.rgb = rgb("7590B2"); tri.line.fill.background()
    circle_label(s, "B&B", 10.31, 3.88, 1.24, TEAL, 18)
    add_text(s, "近似下界", 9.9, 2.92, 2.05, 0.24, 11, "B7C6D9", align=PP_ALIGN.CENTER)
    add_text(s, "可行解 + 最优性证明", 9.73, 5.25, 2.42, 0.28, 11, "B7C6D9", align=PP_ALIGN.CENTER)
    add_text(s, "核心判断", 0.89, 4.46, 1.15, 0.3, 12, "77D5D0", True)
    add_text(s, "LR 仍是重要基线；但当规律无法跨规模稳定迁移时，\n研究主线应转向“可验证、可比较、可学习”的搜索过程。", 0.89, 4.83, 8.0, 0.88, 18, WHITE, True)
    add_text(s, "UC-BnB-LLM · 2026.09", 0.89, 6.79, 4.0, 0.25, 10, "90A4BF")

    # 2 — LR
    s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    title(s, "01 · Motivation", "LR 为什么曾经合理，又为什么遇到瓶颈", 2)
    add_text(s, "LR 的吸引力：把跨机组耦合约束“定价”，将大问题拆成多个单机问题。", 0.76, 1.45, 11.7, 0.38, 17, NAVY, True)
    # process
    steps = [("原始 UC", "负荷 / 备用\n耦合所有机组", PALE_BLUE, BLUE),
             ("松弛耦合约束", "乘子把违约\n转成价格", PALE_ORANGE, ORANGE),
             ("单机子问题", "可并行、结构清晰", PALE_TEAL, TEAL),
             ("恢复可行计划", "修复供需与\n运行限制", PALE_RED, RED)]
    xs = [0.8, 3.88, 6.96, 10.04]
    for idx, ((hd, sub, fill, accent), x) in enumerate(zip(steps, xs), 1):
        rect(s, x, 2.02, 2.42, 1.28, fill=fill, line=accent)
        circle_label(s, str(idx), x + 0.17, 2.18, 0.34, accent, 11)
        add_text(s, hd, x + 0.62, 2.13, 1.62, 0.3, 15, NAVY, True)
        add_text(s, sub, x + 0.2, 2.56, 2.02, 0.52, 12, MUTED, align=PP_ALIGN.CENTER)
        if idx < 4:
            arrow = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x + 2.55), Inches(2.45), Inches(0.46), Inches(0.42))
            arrow.fill.solid(); arrow.fill.fore_color.rgb = rgb(LINE); arrow.line.fill.background()
    # lower split
    rect(s, 0.8, 3.72, 5.72, 2.7, WHITE, LINE)
    add_text(s, "它解决了什么", 1.05, 3.98, 2.0, 0.32, 17, TEAL, True)
    bullet(s, "降低单次迭代的计算难度", 1.08, 4.46, 4.9)
    bullet(s, "自然利用机组可分解结构", 1.08, 4.98, 4.9)
    bullet(s, "快速提供下界与调度启发", 1.08, 5.50, 4.9)
    rect(s, 6.77, 3.72, 5.56, 2.7, WHITE, LINE)
    add_text(s, "它没有自动解决什么", 7.03, 3.98, 2.5, 0.32, 17, RED, True)
    bullet(s, "对偶解不一定是合法 UC 计划", 7.05, 4.46, 4.85, dot=RED)
    bullet(s, "仍需乘子更新与原问题恢复", 7.05, 4.98, 4.85, dot=RED)
    bullet(s, "不能天然给出全局最优证明", 7.05, 5.50, 4.85, dot=RED)
    footer(s, "来源：Gollmer et al., Unit commitment in power generation (2000)；项目 LR 历史经验。")

    # 3 — LR dilemma
    s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    title(s, "02 · Problem", "LR 路线的核心困境：规模放大后，经验不再稳定", 3)
    # scale graphic
    add_text(s, "项目观察", 0.82, 1.48, 1.3, 0.25, 12, TEAL, True)
    add_text(s, "小规模有效 ≠ 大规模仍有效", 0.82, 1.78, 5.3, 0.44, 22, NAVY, True)
    line(s, 1.22, 3.1, 6.1, 3.1, LINE, 4)
    for x, label, sub, col in [(1.1, "50 台", "规律明显", TEAL), (3.35, "200 台", "效果波动", ORANGE), (5.55, "500 台", "可能失效", RED)]:
        circle_label(s, "", x, 2.83, 0.55, col, 1)
        add_text(s, label, x - 0.32, 3.50, 1.2, 0.3, 17, col, True, align=PP_ALIGN.CENTER)
        add_text(s, sub, x - 0.45, 3.88, 1.45, 0.24, 11, MUTED, align=PP_ALIGN.CENTER)
    add_text(s, "同一套乘子更新 / 修复规则", 1.76, 2.54, 3.7, 0.25, 11, MUTED, align=PP_ALIGN.CENTER)
    rect(s, 0.82, 4.56, 5.45, 1.42, PALE_RED, RED)
    add_text(s, "结果：难以区分“规律失效”还是“恢复与调参失效”", 1.08, 4.86, 4.95, 0.72, 17, RED, True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    # cards
    cards = [("对偶间隙", "下界与合法计划之间\n可能长期留有距离"),
             ("原问题恢复", "便宜的对偶安排\n不一定满足供需"),
             ("参数敏感", "步长、乘子和修复\n随实例变化"),
             ("难以归因", "性能变化来自算法、\n恢复还是数据？")]
    positions = [(6.73,1.55),(9.68,1.55),(6.73,4.05),(9.68,4.05)]
    for (hd, sub), (x,y) in zip(cards, positions):
        rect(s, x, y, 2.53, 1.88, WHITE, LINE)
        add_text(s, hd, x+0.22, y+0.24, 2.05, 0.3, 16, NAVY, True)
        line(s, x+0.22, y+0.67, x+0.82, y+0.67, ORANGE, 3)
        add_text(s, sub, x+0.22, y+0.84, 2.06, 0.72, 13, MUTED, align=PP_ALIGN.CENTER)
    footer(s, "判断边界：LR 不是无效，而是当前研究目标需要更稳定的跨规模评价与最优性证据。")

    # 4 — pivot
    s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    title(s, "03 · Decision", "转向分支定界：把“找方案”和“证明方案”放进同一框架", 4)
    # tree
    rect(s, 0.82, 1.56, 5.05, 4.96, WHITE, LINE)
    add_text(s, "分支定界的最小逻辑", 1.08, 1.86, 2.7, 0.34, 17, NAVY, True)
    circle_label(s, "LP", 2.98, 2.42, 0.64, BLUE, 14)
    line(s, 3.3, 3.06, 2.15, 3.72, BLUE, 2); line(s, 3.3, 3.06, 4.45, 3.72, BLUE, 2)
    circle_label(s, "0", 1.84, 3.62, 0.62, TEAL, 15); circle_label(s, "1", 4.14, 3.62, 0.62, TEAL, 15)
    line(s, 2.15, 4.24, 1.48, 4.84, "8CA4C3", 1.5); line(s, 2.15, 4.24, 2.72, 4.84, "8CA4C3", 1.5)
    line(s, 4.45, 4.24, 3.88, 4.84, "8CA4C3", 1.5); line(s, 4.45, 4.24, 5.08, 4.84, "8CA4C3", 1.5)
    for x in [1.27,2.51,3.67,4.87]: circle_label(s, "", x, 4.75, 0.4, "8CA4C3", 1)
    add_text(s, "分支：把一个小数开关拆成 0 / 1", 1.26, 5.52, 4.18, 0.3, 14, NAVY, True, align=PP_ALIGN.CENTER)
    add_text(s, "定界：用下界批量排除不可能更优的子树", 1.14, 5.93, 4.42, 0.3, 13, MUTED, align=PP_ALIGN.CENTER)
    # comparison
    add_text(s, "为什么更符合当前目标", 6.28, 1.56, 3.8, 0.34, 18, NAVY, True)
    rows = [("输出", "LR", "B&B"), ("合法计划", "需额外恢复", "求解器内维护"),
            ("最优性", "通常是下界/近似", "gap 可度量并证明"),
            ("比较对象", "乘子与恢复规则", "节点、时间、gap"),
            ("学习入口", "调乘子/启发式", "学习分支决策")]
    y0=2.05; heights=[0.48,0.72,0.72,0.72,0.72]
    widths=[1.55,2.05,2.2]; xs=[6.28,7.83,9.88]
    for ri,row in enumerate(rows):
        y=y0+sum(heights[:ri]); fill=NAVY if ri==0 else (WHITE if ri%2 else "F0F4F9")
        for ci,(val,x,w) in enumerate(zip(row,xs,widths)):
            rect(s,x,y,w,heights[ri],fill=fill,line=LINE,radius=False)
            add_text(s,val,x+0.08,y+0.05,w-0.16,heights[ri]-0.1,12 if ri else 13,WHITE if ri==0 else TEXT,ri==0 or ci==0,align=PP_ALIGN.CENTER,valign=MSO_ANCHOR.MIDDLE)
    rect(s, 6.28, 6.10, 5.8, 0.46, PALE_TEAL, TEAL)
    add_text(s, "研究焦点：不改 UC 正确答案，只学习“先拆哪一道题”", 6.46, 6.19, 5.44, 0.25, 14, TEAL, True, align=PP_ALIGN.CENTER)
    footer(s, "来源：SCIP branch-and-bound framework；本项目研究边界定义。")

    # 5 — research timeline
    s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    title(s, "04 · Related Work", "学习分支研究：从模仿专家到 LLM 生成可解释程序", 5)
    # timeline
    line(s, 1.1, 2.31, 12.0, 2.31, "B9C5D5", 2.5)
    timeline = [(1.25,"2019","Gasse et al.","二部图 + GNN\n模仿 strong branching",BLUE),
                (4.25,"2020","Ecole","把求解器决策封装为\n可训练、可评测环境",TEAL),
                (7.25,"2020–25","后续研究","泛化、部署开销、\nscore-based policy",ORANGE),
                (10.25,"2026","LLM4Branch","LLM 搜索可执行的\n变量评分程序",RED)]
    for x,yr,hd,sub,col in timeline:
        circle_label(s, "", x, 2.04, 0.54, col, 1)
        add_text(s, yr, x-0.28, 1.47, 1.1, 0.28, 14, col, True, align=PP_ALIGN.CENTER)
        add_text(s, hd, x-0.63, 2.78, 1.8, 0.3, 15, NAVY, True, align=PP_ALIGN.CENTER)
        add_text(s, sub, x-0.73, 3.20, 2.0, 0.7, 12, MUTED, align=PP_ALIGN.CENTER)
    rect(s, 0.86, 4.35, 11.62, 1.86, WHITE, LINE)
    add_text(s, "LLM4Branch 的关键变化", 1.13, 4.62, 2.5, 0.32, 17, NAVY, True)
    items=[("输入", "91 维通用变量特征"), ("搜索", "LLM 修改 Python 评分公式"), ("评价", "真实 SCIP 节点 / 时间 / gap"), ("部署", "只运行固定程序，不逐节点调用 LLM")]
    for idx,(hd,sub) in enumerate(items):
        x=1.12+idx*2.82
        if idx: line(s,x-0.17,4.82,x-0.17,5.89,LINE,1)
        add_text(s,hd,x,5.08,0.7,0.28,12,TEAL,True)
        add_text(s,sub,x,5.40,2.38,0.5,13,TEXT,True if idx==3 else False)
    footer(s, "来源：Gasse et al., NeurIPS 2019, arXiv:1906.01629；Ecole, arXiv:2011.06069；LLM4Branch, arXiv:2605.10401。")

    # 6 — reproduction
    s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    title(s, "05 · Reproduction", "我们的复现已从“小规模可运行”推进到“公开模型可核对”", 6)
    # pipeline
    stages=[("Phase −1","锁定上游环境\n跑通搜索链路"),("Phase 0","UC + 变量评分\nSCIP / Ecole 接通"),("Phase 1–2","结构分支 DSL\n完成小规模试验"),("Public UC","官方数据 + MPS\n进入目标规模")]
    for i,(hd,sub) in enumerate(stages):
        x=0.82+i*3.02
        rect(s,x,1.62,2.55,1.17,PALE_TEAL if i==3 else WHITE,TEAL if i==3 else LINE)
        add_text(s,hd,x+0.18,1.82,2.18,0.27,14,TEAL if i==3 else BLUE,True,align=PP_ALIGN.CENTER)
        add_text(s,sub,x+0.15,2.18,2.25,0.46,12,MUTED,align=PP_ALIGN.CENTER)
        if i<3:
            arr=s.shapes.add_shape(MSO_SHAPE.CHEVRON,Inches(x+2.62),Inches(1.98),Inches(0.30),Inches(0.4)); arr.fill.solid();arr.fill.fore_color.rgb=rgb(LINE);arr.line.fill.background()
    # metrics left
    metrics=[("28","公开 UC 实例"),("12 + 16","OR-LIB / Tejada19"),("24 h","统一时间跨度")]
    for i,(num,label) in enumerate(metrics):
        x=0.86+i*1.85
        add_text(s,num,x,3.28,1.55,0.45,24,NAVY,True,align=PP_ALIGN.CENTER)
        add_text(s,label,x,3.76,1.55,0.3,11,MUTED,align=PP_ALIGN.CENTER)
    rect(s,0.82,4.28,5.62,1.72,WHITE,LINE)
    add_text(s,"数据冻结",1.08,4.52,1.4,0.3,16,NAVY,True)
    bullet(s,"Train 8 · Val 2 · Test 9 · OOD 9",1.12,4.94,4.85,size=14)
    bullet(s,"SHA-256 防止文件与分组悄然变化",1.12,5.38,4.85,size=14)
    # validation right
    rect(s,6.78,3.20,5.55,2.80,WHITE,LINE)
    add_text(s,"官方模型双向核对",7.05,3.47,2.5,0.32,17,NAVY,True)
    rows=[("OR-LIB 200", "35,875,463.302801", "方案固定后一致"),
          ("Tejada 214", "3,250,183.600388", "双方均达最优")]
    for i,(name,obj,res) in enumerate(rows):
        y=4.05+i*0.72
        rect(s,7.04,y,4.98,0.57,"F2F6FA",LINE,radius=False)
        add_text(s,name,7.16,y+0.12,1.25,0.23,12,NAVY,True)
        add_text(s,obj,8.47,y+0.12,1.76,0.23,11,TEXT,align=PP_ALIGN.CENTER)
        add_text(s,res,10.30,y+0.12,1.56,0.23,11,TEAL,True,align=PP_ALIGN.CENTER)
    add_text(s,"关键升级：正式分支实验读取 UnitCommitment.jl 官方 MPS，\n不再依赖手工复刻所有强化约束。",7.05,5.53,4.98,0.45,12,MUTED)
    footer(s, "来源：项目冻结清单、public_uc_crosscheck_evidence.json、UnitCommitment.jl v0.3 官方模型。")

    # 7 — results and next
    s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(BG)
    title(s, "06 · Evidence & Next", "当前证据：分支策略可运行，但大规模瓶颈首先出现在根节点", 7)
    # top stats
    stats=[("16","<500 台官方 MPS"),("44","自适应求解次数"),("42.7 min","累计墙钟"),("3 / 16","分支活跃实例")]
    for i,(num,lab) in enumerate(stats):
        x=0.82+i*3.02
        rect(s,x,1.51,2.55,0.92,WHITE,LINE)
        add_text(s,num,x+0.14,1.67,0.95,0.34,21,TEAL if i==3 else NAVY,True,align=PP_ALIGN.CENTER)
        add_text(s,lab,x+1.02,1.70,1.38,0.28,11,MUTED,align=PP_ALIGN.CENTER)
    # table
    add_text(s,"60 秒结果",0.84,2.74,1.5,0.3,16,NAVY,True)
    headers=["实例","默认 SCIP","RPB","最佳公开程序"]
    widths=[1.20,2.00,1.85,2.15]; xs=[0.84,2.04,4.04,5.89]
    for x,w,h in zip(xs,widths,headers):
        rect(s,x,3.14,w,0.43,NAVY,NAVY,radius=False);add_text(s,h,x+0.04,3.21,w-0.08,0.2,11,WHITE,True,align=PP_ALIGN.CENTER)
    data=[("214","18.6s 最优","56.3s 最优","indset · 4.28e−7"),
          ("250","32.2s 最优","1.97e−6","facilities · 5.65e−7"),
          ("290","30.1s 最优","9.86e−7","facilities · 1.47e−7")]
    for ri,row in enumerate(data):
        y=3.57+ri*0.51; fill=WHITE if ri%2==0 else "F0F4F9"
        for x,w,val in zip(xs,widths,row):
            rect(s,x,y,w,0.51,fill,LINE,radius=False);add_text(s,val,x+0.04,y+0.12,w-0.08,0.22,11,TEXT,ri==0 and x==xs[0],align=PP_ALIGN.CENTER)
    rect(s,0.84,5.25,7.20,0.88,PALE_RED,RED)
    add_text(s,"根节点受限：12 个 OR-LIB + Tejada 480\n在第一次分支前已耗尽 60 秒，评分程序没有控制权。",1.08,5.43,6.72,0.48,13,RED,True,align=PP_ALIGN.CENTER)
    # right conclusions
    rect(s,8.38,2.74,3.94,3.39,WHITE,LINE)
    add_text(s,"路线判断",8.68,3.02,1.5,0.3,17,NAVY,True)
    bullet(s,"默认 SCIP 当前仍是最强实用基准",8.68,3.48,3.30,size=14)
    bullet(s,"91 维特征可复用，公开规则不稳定",8.68,4.04,3.30,size=14)
    bullet(s,"214–290 台：继续学习 UC 专用规则",8.68,4.60,3.30,size=14)
    bullet(s,"≥480 台：先研究根节点与模型表示",8.68,5.16,3.30,size=14,dot=ORANGE)
    rect(s,8.38,6.33,3.94,0.48,PALE_TEAL,TEAL)
    add_text(s,"下一步：双轨推进，而非盲目延长分支时间",8.55,6.43,3.60,0.22,12,TEAL,True,align=PP_ALIGN.CENTER)
    footer(s, "来源：UNDER500_OFFICIAL_MPS_RESULTS_ZH.md；本轮 44 次官方 MPS 自适应实验。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(make_deck())
